import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.pdfgen import canvas

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_word_document(filename):
    doc = docx.Document()
    
    # Page Margins: 0.75 in (54 pt)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    navy = RGBColor(30, 58, 138)       # #1E3A8A
    dark = RGBColor(17, 24, 39)        # #111827
    gray = RGBColor(75, 85, 99)        # #4B5563
    blue = RGBColor(37, 99, 235)       # #2563EB

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.space_before = Pt(0)
    run_title = p_title.add_run("ARGUS PLATFORM")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = navy

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Open-Source Technology Stack & Architectural Specification Report")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = blue
    run_sub.font.bold = True

    # Meta banner
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(16)
    r_meta = p_meta.add_run("Document Version: 1.0  |  Classification: Technical Architecture  |  Zero Cloud / 100% Local Inference")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = gray

    # 1. Executive Summary
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Executive Summary")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = navy

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.space_after = Pt(10)
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.add_run(
        "Argus is an autonomous, real-time biometric trust, anti-spoofing verification, and interview proctoring platform. "
        "It achieves sub-20ms multi-modal verification with zero external API keys and $0 recurring cloud compute costs. "
        "By replacing brittle external cloud LLMs and proprietary facial recognition APIs with hardened open-source neural networks "
        "and client-side accelerated computer vision, Argus guarantees strict data privacy, mathematical bilateral symmetry, and deterministic resilience."
    )

    # 2. Open-Source AI/ML Models
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. Open-Source Computer Vision & AI Models")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = navy

    table_data = [
        ("Component", "Open Source Model", "Framework / Runtime", "Role & Impact"),
        ("Face Detection", "UltraFace Slim 320", "ONNX Runtime (CPU)\n~1.2 MB weights", "Real-time 320x240 multi-scale anchor face detector. Enforces Single-Person Policy via IoU & center-proximity NMS to prevent candidate substitution or multi-person presence."),
        ("Anti-Spoofing (PAD)", "MiniFASNetV2-SE", "ONNX Runtime (CPU)\n~2.7 MB weights", "Feathered multi-scale passive facial liveness classifier. Detects presentation attacks (screens, printed paper, cutouts, video replays, 3D masks) in < 5ms."),
        ("Gaze & Eyelid Tracking", "Google MediaPipe Iris", "Wasm / WebGL (Client)\n478 3D Landmarks", "Client-side iris vector and Eye Aspect Ratio (EAR) tracker running at 60 FPS. Detects looking away (> 30° deviation) while suppressing false alerts during natural blinks."),
        ("Head Pose Proctoring", "Projective 3D Facial Geometry", "Java 21 In-Memory Engine\n< 1ms Execution", "Photometric landmark centroid engine calculating Roll, Pitch, and Yaw. Guarantees mathematical bilateral symmetry for left/right turns with instant > 30° red alert trigger.")
    ]

    t_models = doc.add_table(rows=len(table_data), cols=4)
    t_models.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(1.3), Inches(1.5), Inches(1.6), Inches(2.6)]

    for row_idx, row in enumerate(t_models.rows):
        is_header = (row_idx == 0)
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            if is_header:
                set_cell_background(cell, "1E3A8A")
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(table_data[row_idx][col_idx])
            run.font.name = "Arial"
            if is_header:
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.size = Pt(8.5)
                run.font.color.rgb = dark

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. Technology Stack & Architectural Layers
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. Complete Technology Stack by Layer")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = navy

    tech_layers = [
        ("Backend & Microservice Layer", [
            ("Java 21 LTS: ", "Modern long-term-support runtime providing high throughput, virtual threads, and strongly typed memory safety."),
            ("Spring Boot 3.2.3: ", "Enterprise microservices framework handling REST controllers, static WebRTC assets, JSON marshalling, and actuator telemetry."),
            ("Microsoft ONNX Runtime Java (1.17.1): ", "C++ high-performance native tensor runtime binding executing quantized models directly on CPU with zero GPU requirement."),
            ("Maven Wrapper (mvnw.cmd): ", "Hermetic build pipeline managing dependency trees, automated unit testing, and reproducible packaging.")
        ]),
        ("Frontend & Edge Processing Layer", [
            ("HTML5 & WebRTC: ", "Native web standard accessing local camera feeds at 30–60 FPS with hardware camera driver acceleration."),
            ("Google MediaPipe FaceMesh / Iris: ", "Client-side WebAssembly/WebGL library tracking 478 3D facial landmarks without transmitting raw video across networks."),
            ("HTML5 Canvas 2D & WebGL: ", "Real-time rendering engine drawing ocular tracking arcs, gaze vectors, and pulsating fullscreen red anti-cheat vignettes."),
            ("Vanilla ES6+ JavaScript: ", "Lightweight, zero-framework reactive frontend guaranteeing instantaneous UI response without virtual-DOM overhead.")
        ]),
        ("Mathematical & Algorithmic Core", [
            ("Projective 3D Facial Geometry: ", "Nasal-ocular offset ratio calculation ((noseOffset / (eyeSpan / 2)) * 65.0°) providing lighting-invariant symmetric yaw."),
            ("Eye Aspect Ratio (EAR) Gate: ", "Eyelid height-to-width ratio (EAR < 0.14) that freezes gaze debouncers during resting/blinking to prevent false positives."),
            ("Deterministic Tabular Forensic Engine: ", "Replaced external cloud LLM API with an in-process heuristic scoring engine analyzing cardiac pulse and reaction latencies in < 1ms.")
        ])
    ]

    for category, items in tech_layers:
        p_cat = doc.add_paragraph()
        p_cat.paragraph_format.space_before = Pt(6)
        p_cat.paragraph_format.space_after = Pt(3)
        r_cat = p_cat.add_run(category)
        r_cat.font.name = "Arial"
        r_cat.font.size = Pt(11)
        r_cat.font.bold = True
        r_cat.font.color.rgb = blue

        for bold_prefix, text in items:
            p_item = doc.add_paragraph(style='List Bullet')
            p_item.paragraph_format.space_before = Pt(0)
            p_item.paragraph_format.space_after = Pt(2)
            p_item.paragraph_format.line_spacing = 1.15
            r_b = p_item.add_run(bold_prefix)
            r_b.font.name = "Arial"
            r_b.font.size = Pt(9.5)
            r_b.font.bold = True
            r_b.font.color.rgb = dark
            r_t = p_item.add_run(text)
            r_t.font.name = "Arial"
            r_t.font.size = Pt(9.5)
            r_t.font.color.rgb = dark

    # 4. Open-Source vs Cloud API Comparison
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. Open-Source Architecture vs Cloud APIs")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = navy

    comp_data = [
        ("Architectural Metric", "Proprietary Cloud Solutions", "Argus Open-Source Architecture"),
        ("Operational Cost", "$0.01 - $0.05 per verification frame", "$0.00 (Zero recurring costs)"),
        ("External API Keys", "Required (Gemini, AWS Rekognition, etc.)", "0 API Keys Required"),
        ("Inference Latency", "350ms - 1500ms (Network roundtrips)", "5ms - 20ms (Local on-device / in-memory)"),
        ("Data Privacy & GDPR", "Candidate video streamed to 3rd-party clouds", "100% On-Device / Local. No video leaves machine."),
        ("Offline / Air-Gapped", "Non-functional without active internet", "Fully operational in offline/isolated environments"),
        ("Customizability", "Black-box proprietary models", "Full control over weights, thresholds, and logic")
    ]

    t_comp = doc.add_table(rows=len(comp_data), cols=3)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_widths = [Inches(1.8), Inches(2.5), Inches(2.7)]

    for row_idx, row in enumerate(t_comp.rows):
        is_header = (row_idx == 0)
        for col_idx, cell in enumerate(row.cells):
            cell.width = c_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            if is_header:
                set_cell_background(cell, "1E3A8A")
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(comp_data[row_idx][col_idx])
            run.font.name = "Arial"
            if is_header:
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.size = Pt(8.5)
                run.font.color.rgb = dark

    # 5. Conclusion & Verification Summary
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. Verification & Test Summary")
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = navy

    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.line_spacing = 1.15
    p_conc.add_run(
        "The Argus platform underwent rigorous regression testing across all modules. "
        "The automated test suite (Maven Surefire / JUnit 5) verified 60/60 tests passing with 0 failures. "
        "Key verified capabilities include synthetic face-mirroring symmetry (< 0.005° error), multi-face anchor rejection, "
        "and client-side 60 FPS gaze tracking parity between production static assets and standalone test harnesses."
    )

    doc.save(filename)
    print(f"[SUCCESS] Word document generated: {filename}")


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#6B7280"))
        # Header rule & title
        self.setStrokeColor(colors.HexColor("#E5E7EB"))
        self.setLineWidth(0.5)
        self.line(40, letter[1] - 40, letter[0] - 40, letter[1] - 40)
        self.drawString(40, letter[1] - 35, "Argus Platform - Open-Source Technology Stack Report")
        
        # Footer rule & page number
        self.line(40, 40, letter[0] - 40, 40)
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 40, 28, page_text)
        self.drawString(40, 28, "Confidential - Engineering Architecture Specification")
        self.restoreState()


def create_pdf_document(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=55,
        bottomMargin=55
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=2
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#2563EB"),
        spaceAfter=8
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#4B5563"),
        spaceAfter=12
    )

    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=10,
        spaceAfter=5
    )

    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#2563EB"),
        spaceBefore=5,
        spaceAfter=2
    )

    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=4
    )

    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )

    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1F2937")
    )

    story = []

    # Title & Header
    story.append(Paragraph("ARGUS BIOMETRIC PROCTORING PLATFORM", title_style))
    story.append(Paragraph("Open-Source Technology Stack & Architectural Specification Report", subtitle_style))
    story.append(Paragraph("Version: 1.0  |  Classification: Technical Architecture  |  Zero Cloud / 100% Local Inference", meta_style))

    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary", h1_style))
    story.append(Paragraph(
        "Argus is an autonomous, on-device biometric trust, anti-spoofing verification, and interview proctoring platform. "
        "It achieves sub-20ms multi-modal verification with zero external API keys and $0 recurring cloud compute costs. "
        "By replacing brittle external cloud LLMs and proprietary facial recognition APIs with hardened open-source neural networks "
        "and client-side accelerated computer vision, Argus guarantees strict data privacy, mathematical bilateral symmetry, and deterministic resilience.",
        body_style
    ))

    # 2. Open-Source AI/ML Models
    story.append(Paragraph("2. Open-Source Computer Vision & AI Models", h1_style))
    
    model_table_data = [
        [
            Paragraph("Component", table_header_style),
            Paragraph("Open-Source Model", table_header_style),
            Paragraph("Framework / Runtime", table_header_style),
            Paragraph("Role & Architectural Impact", table_header_style)
        ],
        [
            Paragraph("<b>Face Detection</b>", table_cell_style),
            Paragraph("UltraFace Slim 320", table_cell_style),
            Paragraph("ONNX Runtime (CPU)<br/>~1.2 MB weights", table_cell_style),
            Paragraph("Real-time 320x240 multi-scale anchor face detector. Enforces Single-Person Policy via IoU and center-proximity NMS to prevent candidate substitution or multi-person presence.", table_cell_style)
        ],
        [
            Paragraph("<b>Anti-Spoofing (PAD)</b>", table_cell_style),
            Paragraph("MiniFASNetV2-SE", table_cell_style),
            Paragraph("ONNX Runtime (CPU)<br/>~2.7 MB weights", table_cell_style),
            Paragraph("Feathered multi-scale passive facial liveness classifier. Detects presentation attacks (screens, printed paper, cutouts, video replays, 3D masks) in < 5ms.", table_cell_style)
        ],
        [
            Paragraph("<b>Gaze & Eye Tracking</b>", table_cell_style),
            Paragraph("Google MediaPipe Iris", table_cell_style),
            Paragraph("Wasm / WebGL (Client)<br/>478 3D Landmarks", table_cell_style),
            Paragraph("Client-side iris vector and Eye Aspect Ratio (EAR) tracker running at 60 FPS. Detects looking away (> 30° deviation) while suppressing false alerts during natural blinks.", table_cell_style)
        ],
        [
            Paragraph("<b>Head Movement Proctoring</b>", table_cell_style),
            Paragraph("Projective 3D Facial Geometry", table_cell_style),
            Paragraph("Java 21 In-Memory Engine<br/>< 1ms Execution", table_cell_style),
            Paragraph("Photometric landmark centroid engine calculating Roll, Pitch, and Yaw. Guarantees mathematical bilateral symmetry for left/right turns with instant > 30° red alert trigger.", table_cell_style)
        ]
    ]

    t_mod = Table(model_table_data, colWidths=[100, 110, 110, 212])
    t_mod.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_mod)

    # 3. Complete Technology Stack by Layer
    story.append(Paragraph("3. Complete Technology Stack by Architectural Layer", h1_style))
    
    layers = [
        ("Backend & Microservice Layer", [
            ("Java 21 LTS: ", "Modern long-term-support runtime providing high throughput, virtual threads, and strongly typed memory safety."),
            ("Spring Boot 3.2.3: ", "Enterprise microservices framework handling REST controllers, static WebRTC assets, and JSON marshalling."),
            ("Microsoft ONNX Runtime Java (1.17.1): ", "C++ high-performance native tensor runtime binding executing quantized models directly on CPU with 0 GPU requirement."),
            ("Maven Wrapper (mvnw.cmd): ", "Hermetic build pipeline managing dependency trees, automated unit testing, and reproducible packaging.")
        ]),
        ("Frontend & Edge Processing Layer", [
            ("HTML5 & WebRTC: ", "Native web standard accessing local camera feeds at 30–60 FPS with hardware camera driver acceleration."),
            ("Google MediaPipe FaceMesh / Iris: ", "Client-side WebAssembly/WebGL library tracking 478 3D facial landmarks without transmitting raw video across networks."),
            ("HTML5 Canvas 2D & WebGL: ", "Real-time rendering engine drawing ocular tracking arcs, gaze vectors, and pulsating fullscreen red anti-cheat vignettes."),
            ("Vanilla ES6+ JavaScript: ", "Lightweight, zero-framework reactive frontend guaranteeing instantaneous UI response without virtual-DOM overhead.")
        ]),
        ("Mathematical & Algorithmic Core", [
            ("Projective 3D Facial Geometry: ", "Nasal-ocular offset ratio calculation ((noseOffset / (eyeSpan / 2)) * 65.0°) providing lighting-invariant symmetric yaw."),
            ("Eye Aspect Ratio (EAR) Gate: ", "Eyelid height-to-width ratio (EAR < 0.14) that freezes gaze debouncers during resting/blinking to prevent false positives."),
            ("Deterministic Tabular Forensic Engine: ", "Replaced external cloud LLM API with an in-process heuristic scoring engine analyzing cardiac pulse and reaction latencies in < 1ms.")
        ])
    ]

    for cat_name, items in layers:
        story.append(Paragraph(cat_name, h2_style))
        for prefix, body in items:
            bullet_text = f"• <b>{prefix}</b> {body}"
            story.append(Paragraph(bullet_text, body_style))

    # 4. Open-Source vs Cloud API Comparison Table
    story.append(Paragraph("4. Open-Source Architecture vs Cloud APIs", h1_style))
    comp_table_data = [
        [
            Paragraph("Architectural Metric", table_header_style),
            Paragraph("Proprietary Cloud Solutions", table_header_style),
            Paragraph("Argus Open-Source Architecture", table_header_style)
        ],
        [
            Paragraph("<b>Operational Cost</b>", table_cell_style),
            Paragraph("$0.01 - $0.05 per verification frame", table_cell_style),
            Paragraph("<b>$0.00</b> (Zero recurring costs)", table_cell_style)
        ],
        [
            Paragraph("<b>External API Keys</b>", table_cell_style),
            Paragraph("Required (Gemini, AWS Rekognition, etc.)", table_cell_style),
            Paragraph("<b>0 API Keys Required</b>", table_cell_style)
        ],
        [
            Paragraph("<b>Inference Latency</b>", table_cell_style),
            Paragraph("350ms - 1500ms (Network roundtrips)", table_cell_style),
            Paragraph("<b>5ms - 20ms</b> (Local on-device / in-memory)", table_cell_style)
        ],
        [
            Paragraph("<b>Data Privacy & GDPR</b>", table_cell_style),
            Paragraph("Candidate video streamed to 3rd-party clouds", table_cell_style),
            Paragraph("<b>100% On-Device / Local</b>. No video leaves machine.", table_cell_style)
        ],
        [
            Paragraph("<b>Offline / Air-Gapped</b>", table_cell_style),
            Paragraph("Non-functional without active internet", table_cell_style),
            Paragraph("<b>Fully operational</b> in offline/isolated environments", table_cell_style)
        ],
        [
            Paragraph("<b>Customizability</b>", table_cell_style),
            Paragraph("Black-box proprietary models", table_cell_style),
            Paragraph("<b>Full control</b> over weights, thresholds, and logic", table_cell_style)
        ]
    ]

    t_comp = Table(comp_table_data, colWidths=[120, 190, 222])
    t_comp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
    ]))
    story.append(t_comp)

    # 5. Verification & Test Summary
    story.append(Paragraph("5. Verification & Test Summary", h1_style))
    story.append(Paragraph(
        "The Argus platform underwent rigorous regression testing across all modules. "
        "The automated test suite (Maven Surefire / JUnit 5) verified <b>60/60 tests passing</b> with 0 failures. "
        "Key verified capabilities include synthetic face-mirroring symmetry (< 0.005° error), multi-face anchor rejection, "
        "and client-side 60 FPS gaze tracking parity between production static assets and standalone test harnesses.",
        body_style
    ))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[SUCCESS] PDF document generated: {filename}")


if __name__ == "__main__":
    docx_path = os.path.join(os.path.dirname(__file__), "Argus_Tech_Stack_and_Open_Source_Report.docx")
    pdf_path = os.path.join(os.path.dirname(__file__), "Argus_Tech_Stack_and_Open_Source_Report.pdf")
    
    create_word_document(docx_path)
    create_pdf_document(pdf_path)
